"""DOCX export — полный протокол встречи в формате Word."""

from __future__ import annotations
import io
from datetime import datetime
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


# Цвета спикеров (циклически)
SPEAKER_COLORS = [
    RGBColor(0x1A, 0x73, 0xE8),  # blue
    RGBColor(0xD9, 0x34, 0x25),  # red
    RGBColor(0x18, 0x8E, 0x38),  # green
    RGBColor(0xF2, 0x9D, 0x00),  # orange
    RGBColor(0x7B, 0x2F, 0xBE),  # purple
    RGBColor(0x00, 0x7B, 0x83),  # teal
]


def _speaker_color(speaker: str, cache: dict) -> RGBColor:
    if speaker not in cache:
        cache[speaker] = SPEAKER_COLORS[len(cache) % len(SPEAKER_COLORS)]
    return cache[speaker]


def _fmt_time(seconds: float) -> str:
    h = int(seconds // 3600)
    m = int((seconds % 3600) // 60)
    s = int(seconds % 60)
    return f"{h:02d}:{m:02d}:{s:02d}" if h else f"{m:02d}:{s:02d}"


def to_docx(
    segments: list[dict],
    transcript: str = "",
    summary: str = "",
    highlights: str = "",
    chapters: str = "",
    action_items: str = "",
    keywords: list[dict] | None = None,
    speaker_stats: list[dict] | None = None,
    event_title: str = "Протокол встречи",
) -> bytes:
    doc = Document()

    # ── Стили ──────────────────────────────────────────────────
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # ── Заголовок ──────────────────────────────────────────────
    title = doc.add_heading(event_title, 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(f"Дата: {datetime.now().strftime('%d.%m.%Y %H:%M')}")
    doc.add_paragraph()

    # ── Резюме ─────────────────────────────────────────────────
    if summary:
        doc.add_heading("Резюме", 1)
        doc.add_paragraph(summary)
        doc.add_paragraph()

    # ── Highlights ─────────────────────────────────────────────
    if highlights:
        doc.add_heading("Ключевые моменты", 1)
        doc.add_paragraph(highlights)
        doc.add_paragraph()

    # ── Action Items ───────────────────────────────────────────
    if action_items:
        doc.add_heading("Задачи", 1)
        doc.add_paragraph(action_items)
        doc.add_paragraph()

    # ── Статистика спикеров ────────────────────────────────────
    if speaker_stats:
        doc.add_heading("Статистика участников", 1)
        table = doc.add_table(rows=1, cols=5)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        for i, h in enumerate(["Спикер", "Время (сек)", "Сегменты", "Слова", "Доля %"]):
            hdr[i].text = h

        color_cache: dict = {}
        for stat in speaker_stats:
            row = table.add_row().cells
            row[0].text = stat.get("speaker", "")
            row[1].text = str(stat.get("total_seconds", ""))
            row[2].text = str(stat.get("segment_count", ""))
            row[3].text = str(stat.get("word_count", ""))
            row[4].text = f"{stat.get('percentage', 0)}%"

        doc.add_paragraph()

    # ── Главы ──────────────────────────────────────────────────
    if chapters:
        doc.add_heading("Главы", 1)
        doc.add_paragraph(chapters)
        doc.add_paragraph()

    # ── Ключевые слова ─────────────────────────────────────────
    if keywords:
        doc.add_heading("Ключевые слова", 1)
        kw_text = ", ".join(k["keyword"] for k in keywords[:20])
        doc.add_paragraph(kw_text)
        doc.add_paragraph()

    # ── Транскрипция по спикерам ───────────────────────────────
    doc.add_page_break()
    doc.add_heading("Транскрипция", 1)

    color_cache: dict = {}
    current_speaker = None

    for seg in segments:
        speaker   = seg.get("speaker")
        text      = seg.get("text", "").strip()
        start     = _fmt_time(seg.get("start", 0))

        if not text:
            continue

        if speaker and speaker != current_speaker:
            current_speaker = speaker
            p = doc.add_paragraph()
            run = p.add_run(f"[{start}] {speaker}:")
            run.bold = True
            run.font.color.rgb = _speaker_color(speaker, color_cache)

        p = doc.add_paragraph(f"[{start}] {text}" if not speaker else f"        {text}")

    # ── Полный текст транскрипции ──────────────────────────────
    if transcript:
        doc.add_page_break()
        doc.add_heading("Полный текст", 1)
        doc.add_paragraph(transcript)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
